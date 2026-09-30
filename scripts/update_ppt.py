"""외국인 이용객 트렌드 수치를 계산해 기존 PPT의 표 값만 자동으로 갱신한다 - 실습 v1

data/bus_foreign_sample.md 를 읽어 기간별(올해 누계/최근 1년/최근 3년)·노선별
이용객 수와 증감률을 계산하고, ppt/고속시외버스_외국인이용객_트렌드.pptx 의 표
셀 값만 찾아서 바꾼다. 표의 라벨(첫 열)과 헤더, 서식은 그대로 두고 숫자 칸만
교체한다("기존 문구의 수치만 바꾼다"는 원 기획서 취지 유지).

기간 정의(원 기획서에 "확인 필요"로 남아 있던 부분 -> 이 스크립트의 잠정 기준):
- 올해 누계: 최신월이 속한 해의 1월~최신월 합계, 비교: 전년 동기(같은 1월~최신월)
- 최근 1년: 최신월 포함 직전 12개월 합계, 비교: 그 앞 12개월("그 전 1년")
- 최근 3년: 보유 데이터(최근 36개월) 중 가장 최근 12개월 합계,
            비교: 데이터가 시작하는 시점의 12개월("3년 전" 구간)
  (최근 1년과 최근 3년의 "현재값"이 같은 최근 12개월인 것은 원 기획서 7번의
   가상 샘플 표에서도 두 행의 기준값이 125로 같았던 패턴을 그대로 따른 것.)

2장 세부 표 헤더는 "(지수)"로 돼 있으나, 지수화 방식이 미확정이라 이번 버전은
실제 집계 이용객 수(카드+OTA 합산)를 그대로 넣는다. 이 부분은 사람 확인 항목으로
로그에 남긴다.
"""

import re
import sys
from pathlib import Path

from pptx import Presentation

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
PPT_PATH = BASE_DIR / "ppt" / "고속시외버스_외국인이용객_트렌드.pptx"
OUTPUT_DIR = BASE_DIR / "output"

BUS_FILE = "bus_foreign_sample.md"
KOSIS_FILE = "kosis_foreign_sample.md"


def parse_table(path: Path) -> list[dict]:
    if not path.exists():
        sys.exit(f"[오류] 파일을 찾을 수 없습니다: {path}")
    rows = []
    header = None
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if re.fullmatch(r"-+", cells[0]):
            continue
        if header is None:
            header = cells
            continue
        rows.append(dict(zip(header, cells)))
    return rows


def load_bus_rows() -> list[dict]:
    raw = parse_table(DATA_DIR / BUS_FILE)
    rows = []
    for r in raw:
        rows.append(
            {
                "ym": r["연월"],
                "route": r["노선"],
                "total": int(r["해외카드 인원"]) + int(r["해외OTA 인원"]),
            }
        )
    return rows


def shift_month(ym: str, delta: int) -> str:
    y, m = (int(x) for x in ym.split("-"))
    idx = y * 12 + (m - 1) + delta
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def sum_range(rows: list[dict], start_ym: str, end_ym: str, route: str | None = None) -> int:
    total = 0
    for r in rows:
        if start_ym <= r["ym"] <= end_ym and (route is None or r["route"] == route):
            total += r["total"]
    return total


def pct_change(cur: int, base: int) -> float:
    if base == 0:
        return 0.0
    return (cur - base) / base * 100


def fmt_pct(v: float) -> str:
    sign = "+" if v >= 0 else ""
    return f"{sign}{v:.1f}%"


def fmt_num(v: int) -> str:
    return f"{v:,}"


def build_periods(rows: list[dict]) -> dict:
    latest_ym = max(r["ym"] for r in rows)
    earliest_ym = min(r["ym"] for r in rows)
    latest_y, latest_m = (int(x) for x in latest_ym.split("-"))

    this_year_start = f"{latest_y:04d}-01"
    prev_year_start = f"{latest_y - 1:04d}-01"
    prev_year_end = f"{latest_y - 1:04d}-{latest_m:02d}"

    last12_start = shift_month(latest_ym, -11)
    prior12_end = shift_month(last12_start, -1)
    prior12_start = shift_month(prior12_end, -11)

    threeyr_recent_start = shift_month(latest_ym, -11)
    threeyr_base_start = earliest_ym
    threeyr_base_end = shift_month(earliest_ym, 11)

    return {
        "latest_ym": latest_ym,
        "올해 누계": {
            "cur": (this_year_start, latest_ym),
            "base": (prev_year_start, prev_year_end),
            "short": "전년 동기",
        },
        "최근 1년": {
            "cur": (last12_start, latest_ym),
            "base": (prior12_start, prior12_end),
            "short": "그 전 1년",
        },
        "최근 3년": {
            "cur": (threeyr_recent_start, latest_ym),
            "base": (threeyr_base_start, threeyr_base_end),
            "short": "3년 전",
        },
    }


def build_summary_rows(rows: list[dict], periods: dict) -> list[dict]:
    result = []
    for label in ("올해 누계", "최근 1년", "최근 3년"):
        cur_s, cur_e = periods[label]["cur"]
        base_s, base_e = periods[label]["base"]
        cur_total = sum_range(rows, cur_s, cur_e)
        base_total = sum_range(rows, base_s, base_e)
        result.append(
            {
                "label": label,
                "cur": cur_total,
                "cur_range": (cur_s, cur_e),
                "base": base_total,
                "base_range": (base_s, base_e),
                "base_short": periods[label]["short"],
                "change": pct_change(cur_total, base_total),
            }
        )
    return result


def build_detail_rows(rows: list[dict], periods: dict) -> list[dict]:
    routes = sorted({r["route"] for r in rows})
    cur_s, cur_e = periods["올해 누계"]["cur"]
    base_s, base_e = periods["올해 누계"]["base"]
    result = []
    for route in routes:
        cur_total = sum_range(rows, cur_s, cur_e, route=route)
        base_total = sum_range(rows, base_s, base_e, route=route)
        result.append(
            {
                "label": route,
                "cur": cur_total,
                "base": base_total,
                "change": pct_change(cur_total, base_total),
            }
        )
    return result


def find_table(slide):
    for shape in slide.shapes:
        if shape.has_table:
            return shape.table
    return None


def set_cell_text(cell, text: str) -> None:
    """서식(폰트 크기·색 등)은 유지하고 첫 번째 run의 텍스트만 바꾼다."""
    paragraph = cell.text_frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for extra_run in paragraph.runs[1:]:
            extra_run.text = ""
    else:
        cell.text = text


def update_table_by_label(table, values_by_label: dict, value_cols: list[tuple]) -> set:
    """표 1열(라벨)로 행을 찾아 나머지 열의 숫자 칸만 새 값으로 교체."""
    updated = set()
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label not in values_by_label:
            continue
        data = values_by_label[label]
        for col_idx, render in value_cols:
            set_cell_text(row.cells[col_idx], render(data))
        updated.add(label)
    return updated


def main():
    if not PPT_PATH.exists():
        sys.exit(f"[오류] 대상 PPT를 찾을 수 없습니다: {PPT_PATH} (먼저 gen_sample_ppt.py 실행 필요)")

    bus_rows = load_bus_rows()
    periods = build_periods(bus_rows)
    summary = build_summary_rows(bus_rows, periods)
    detail = build_detail_rows(bus_rows, periods)

    prs = Presentation(PPT_PATH)
    if len(prs.slides) < 2:
        sys.exit("[오류] PPT에 슬라이드가 2장 미만입니다.")
    slide1, slide2 = prs.slides[0], prs.slides[1]

    table1 = find_table(slide1)
    table2 = find_table(slide2)
    if table1 is None or table2 is None:
        sys.exit("[오류] PPT에서 표를 찾을 수 없습니다.")

    summary_by_label = {r["label"]: r for r in summary}
    updated1 = update_table_by_label(
        table1,
        summary_by_label,
        value_cols=[
            (1, lambda d: fmt_num(d["cur"])),
            (2, lambda d: f"{d['base_short']} {fmt_num(d['base'])}"),
            (3, lambda d: fmt_pct(d["change"])),
        ],
    )

    detail_by_label = {r["label"]: r for r in detail}
    updated2 = update_table_by_label(
        table2,
        detail_by_label,
        value_cols=[
            (1, lambda d: fmt_num(d["cur"])),
            (2, lambda d: fmt_num(d["base"])),
            (3, lambda d: fmt_pct(d["change"])),
        ],
    )

    missing1 = set(summary_by_label) - updated1
    missing2 = set(detail_by_label) - updated2
    if missing1 or missing2:
        print(f"[경고] PPT 표에서 라벨을 못 찾아 갱신 안 됨: 1장 {missing1 or '-'}, 2장 {missing2 or '-'}")

    prs.save(PPT_PATH)

    OUTPUT_DIR.mkdir(exist_ok=True)
    log_path = OUTPUT_DIR / f"수치교체_로그_{periods['latest_ym']}.md"
    lines = [
        f"# 외국인 이용객 PPT 자동 갱신 로그 (기준 시점: {periods['latest_ym']})",
        "",
        "> 이 로그는 사람이 원본과 대조하기 위한 기록입니다. PPT는 이미 자동 갱신되었습니다.",
        "> 데이터는 가상 샘플이며 실제 수치가 아닙니다.",
        "",
        "## 1장 요약 (전체 노선 합계)",
        "| 구분 | 기준(이용객 수) | 비교 시점 | 비교값 | 증감률 |",
        "|---|---|---|---|---|",
    ]
    for r in summary:
        cur_s, cur_e = r["cur_range"]
        base_s, base_e = r["base_range"]
        lines.append(
            f"| {r['label']} ({cur_s}~{cur_e}) | {fmt_num(r['cur'])} | "
            f"{r['base_short']} ({base_s}~{base_e}) | {fmt_num(r['base'])} | {fmt_pct(r['change'])} |"
        )
    lines += [
        "",
        "## 2장 세부 (노선별, 올해 누계 vs 전년 동기)",
        "| 노선 | 올해 누계 | 전년 동기 | 증감률 |",
        "|---|---|---|---|",
    ]
    for r in detail:
        lines.append(f"| {r['label']} | {fmt_num(r['cur'])} | {fmt_num(r['base'])} | {fmt_pct(r['change'])} |")
    lines += [
        "",
        "## 출처",
        f"- 당사 시스템(가상 샘플, 해외카드+해외OTA 단순 합산): data/{BUS_FILE}",
        f"- 통계청 KOSIS(가상 샘플, 참고용 · 이번 버전에서는 PPT 수치 계산에는 미사용): data/{KOSIS_FILE}",
        "",
        "## 사람 확인 필요",
        "- PPT에 반영된 수치가 위 표와 원본 데이터를 정확히 대조했는지 확인",
        "- '최근 3년' 비교 기준(데이터 시작 12개월 vs 최근 12개월) 정의가 실제 보고 관행과 맞는지 확인",
        "- 해외카드 인원과 해외OTA 인원의 중복 집계 여부 확인(현재는 단순 합산)",
        "- 2장 헤더의 '(지수)' 표기와 달리 이번 버전은 실제 이용객 수를 그대로 넣었음 — 지수화가 필요하면 별도 확인",
        "- PPT 문구(제목·부제)가 갱신된 수치와 자연스럽게 맞는지 확인",
    ]
    log_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(f"PPT 갱신 완료 -> {PPT_PATH}")
    print(f"로그 저장 -> {log_path}")
    for r in summary:
        print(f"  [1장] {r['label']}: {fmt_num(r['cur'])} ({fmt_pct(r['change'])})")
    for r in detail:
        print(f"  [2장] {r['label']}: {fmt_num(r['cur'])} ({fmt_pct(r['change'])})")


if __name__ == "__main__":
    main()
