"""월간 운송대금 KPI 검증 자동화 - 실습 v1

기초데이터(시스템별 원본 3개)와 업데이트 데이터(관리시트)의 금액 합계를 비교해
검증 결과 요약 1장(.xlsx)을 만든다. 숫자 계산과 일치 판정만 다루며, 문맥적 요약
다듬기는 이 결과물을 보고 사람 또는 AI가 이어서 한다.
"""

import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

# (기초데이터 파일명, 시스템 표시명)
BASELINE_FILES = [
    ("system_a_baseline.md", "시스템A"),
    ("system_b_baseline.md", "시스템B"),
    ("system_c_baseline.md", "시스템C"),
]
UPDATED_FILE = "management_sheet_updated.md"


def parse_markdown_table(path: Path) -> list[tuple[str, int]]:
    if not path.exists():
        sys.exit(f"[오류] 파일을 찾을 수 없습니다: {path}")

    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) != 2:
            continue
        label, amount_text = cells
        if re.fullmatch(r"-+", amount_text) or label in ("항목", "시스템"):
            continue  # 헤더/구분선 행 건너뜀
        amount_digits = re.sub(r"[^\d-]", "", amount_text)
        if not amount_digits:
            sys.exit(f"[오류] 금액을 숫자로 읽을 수 없습니다: '{line}' ({path})")
        rows.append((label, int(amount_digits)))
    return rows


def load_baseline_totals() -> dict[str, int]:
    totals = {}
    for filename, system_name in BASELINE_FILES:
        rows = parse_markdown_table(DATA_DIR / filename)
        totals[system_name] = sum(amount for _, amount in rows)
    return totals


def load_updated_totals() -> dict[str, int]:
    rows = parse_markdown_table(DATA_DIR / UPDATED_FILE)
    return dict(rows)


def compare(baseline: dict[str, int], updated: dict[str, int]) -> list[dict]:
    results = []
    for _, system_name in BASELINE_FILES:
        base_total = baseline.get(system_name)
        upd_total = updated.get(system_name)
        if upd_total is None:
            sys.exit(f"[오류] 관리시트에 '{system_name}' 금액이 없습니다.")
        diff = upd_total - base_total
        results.append(
            {
                "system": system_name,
                "baseline": base_total,
                "updated": upd_total,
                "diff": diff,
                "match": diff == 0,
            }
        )

    grand_base = sum(r["baseline"] for r in results)
    grand_upd = sum(r["updated"] for r in results)
    results.append(
        {
            "system": "전체 합계",
            "baseline": grand_base,
            "updated": grand_upd,
            "diff": grand_upd - grand_base,
            "match": grand_upd == grand_base,
        }
    )
    return results


HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MATCH_FILL = PatternFill("solid", fgColor="E2EFDA")
MISMATCH_FILL = PatternFill("solid", fgColor="FCE4D6")
WON_FORMAT = '#,##0"원"'


def render_summary_xlsx(results: list[dict], run_date: datetime) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "검증결과요약"

    ws["A1"] = f"운송대금 KPI 검증 결과 요약 ({run_date.strftime('%Y-%m')})"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")

    ws["A2"] = "실행 일시"
    ws["B2"] = run_date.strftime("%Y-%m-%d %H:%M")
    ws["A3"] = "기초데이터"
    ws["B3"] = "data/system_a_baseline.md, system_b_baseline.md, system_c_baseline.md"
    ws["A4"] = "업데이트 데이터"
    ws["B4"] = "data/management_sheet_updated.md"
    for row in (2, 3, 4):
        ws[f"A{row}"].font = Font(bold=True)

    ws["A6"] = "1. 시스템별 합계 비교"
    ws["A6"].font = Font(bold=True, size=12)

    headers = ["시스템", "기초데이터 합계", "업데이트 데이터 합계", "차이", "판정"]
    header_row = 7
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    mismatches = []
    row = header_row
    for r in results:
        row += 1
        verdict = "일치" if r["match"] else "불일치"
        ws.cell(row=row, column=1, value=r["system"])
        for col, key in ((2, "baseline"), (3, "updated"), (4, "diff")):
            cell = ws.cell(row=row, column=col, value=r[key])
            cell.number_format = WON_FORMAT
        ws.cell(row=row, column=5, value=verdict).alignment = Alignment(horizontal="center")

        fill = MATCH_FILL if r["match"] else MISMATCH_FILL
        for col in range(1, 6):
            ws.cell(row=row, column=col).fill = fill

        if not r["match"] and r["system"] != "전체 합계":
            mismatches.append(r)

    note_row = row + 2
    ws.cell(row=note_row, column=1, value="2. 사람이 확인할 항목").font = Font(bold=True, size=12)
    note_row += 1
    if mismatches:
        for r in mismatches:
            ws.cell(
                row=note_row,
                column=1,
                value=(
                    f"{r['system']}: 기초데이터({r['baseline']:,}원)와 "
                    f"업데이트 데이터({r['updated']:,}원)가 {abs(r['diff']):,}원 차이납니다. "
                    "관리시트 입력 원본과 재대조가 필요합니다."
                ),
            )
            note_row += 1
    else:
        ws.cell(row=note_row, column=1, value="없음 (모든 시스템 수치 일치)")

    column_widths = [14, 20, 22, 16, 10]
    for col, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    return wb


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)

    baseline = load_baseline_totals()
    updated = load_updated_totals()
    results = compare(baseline, updated)

    run_date = datetime.now()
    wb = render_summary_xlsx(results, run_date)

    output_path = OUTPUT_DIR / f"검증결과요약_{run_date.strftime('%Y-%m')}.xlsx"
    wb.save(output_path)

    print(f"검증 완료 -> {output_path}")
    for r in results:
        verdict = "일치" if r["match"] else "불일치"
        print(f"  {r['system']}: {verdict} (차이 {r['diff']:,}원)")


if __name__ == "__main__":
    main()
